"""ATS-safe DOCX export. Visual two-column layouts are flattened."""

from __future__ import annotations

import io
import re
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from app.services.resume_text import format_date_range


def _safe_filename(personal: dict[str, Any], ext: str) -> str:
    name = str(personal.get("fullName") or "Resume").strip()
    slug = re.sub(r"[^A-Za-z0-9]+", "_", name).strip("_") or "Resume"
    return f"{slug}_Resume.{ext}"


def build_docx(data: dict[str, Any], settings: dict[str, Any] | None = None) -> tuple[bytes, str]:
    settings = settings or {}
    document = Document()
    section = document.sections[0]
    margin = {"narrow": 0.5, "normal": 0.7, "wide": 1.0}.get(settings.get("margin", "normal"), 0.7)
    for edge in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
        setattr(section, edge, Inches(margin))

    personal = data.get("personal") or {}
    name = personal.get("fullName") or "Resume"
    heading = document.add_paragraph()
    run = heading.add_run(name)
    run.bold = True
    run.font.size = Pt(18)
    run.font.color.rgb = RGBColor(17, 24, 39)
    heading.alignment = WD_ALIGN_PARAGRAPH.LEFT

    title = personal.get("title")
    if title:
        p = document.add_paragraph(title)
        p.runs[0].italic = True

    contact = [
        personal.get("email"),
        personal.get("phone"),
        personal.get("location"),
        personal.get("linkedin"),
        personal.get("github"),
        personal.get("website"),
        personal.get("portfolio"),
    ]
    contact_line = "  ·  ".join(part for part in contact if part)
    if contact_line:
        document.add_paragraph(contact_line)

    date_style = settings.get("dateFormat") or settings.get("date_format") or "short-month"

    for block in data.get("sections") or []:
        if not block.get("visible", True):
            continue
        kind = block.get("kind")
        title_text = str(block.get("title") or "").upper()
        if not _section_has_content(block):
            continue
        heading_p = document.add_paragraph()
        run = heading_p.add_run(title_text)
        run.bold = True
        run.font.size = Pt(12)
        pPr = heading_p._p.get_or_add_pPr()
        pBdr = pPr.makeelement(qn("w:pBdr"), {})
        bottom = pBdr.makeelement(
            qn("w:bottom"),
            {qn("w:val"): "single", qn("w:sz"): "6", qn("w:space"): "1", qn("w:color"): "111827"},
        )
        pBdr.append(bottom)
        pPr.append(pBdr)

        if kind == "text" and block.get("content"):
            document.add_paragraph(str(block["content"]))
        elif kind == "skills":
            for group in block.get("groups") or []:
                skills = ", ".join(group.get("skills") or [])
                if not skills:
                    continue
                label = group.get("name") or "Skills"
                document.add_paragraph(f"{label}: {skills}")
        elif kind == "tags":
            tags = ", ".join(block.get("tags") or [])
            if tags:
                document.add_paragraph(tags)
        else:
            for item in block.get("items") or []:
                _write_item(document, item, kind, date_style)

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue(), _safe_filename(personal, "docx")


def _write_item(document: Document, item: dict[str, Any], kind: str, date_style: str) -> None:
    if kind == "experience":
        head = " · ".join(part for part in (item.get("title"), item.get("company")) if part)
        dates = format_date_range(
            str(item.get("startDate") or ""),
            str(item.get("endDate") or ""),
            bool(item.get("current")),
            date_style,
        )
        line = document.add_paragraph()
        run = line.add_run(head)
        run.bold = True
        if dates:
            line.add_run(f"    {dates}")
        if item.get("location"):
            document.add_paragraph(str(item["location"]))
        if item.get("description"):
            document.add_paragraph(str(item["description"]))
        for bullet in item.get("bullets") or []:
            document.add_paragraph(str(bullet), style="List Bullet")
        return

    if kind == "education":
        head = " ".join(part for part in (item.get("degree"), item.get("field")) if part)
        line = document.add_paragraph()
        run = line.add_run(head or item.get("institution") or "Education")
        run.bold = True
        if item.get("institution"):
            document.add_paragraph(str(item["institution"]))
        dates = format_date_range(
            str(item.get("startDate") or ""),
            str(item.get("endDate") or ""),
            bool(item.get("current")),
            date_style,
        )
        extras = [dates, item.get("gpa") and f"GPA {item['gpa']}"]
        extra = "  ·  ".join(part for part in extras if part)
        if extra:
            document.add_paragraph(extra)
        for bullet in item.get("bullets") or []:
            document.add_paragraph(str(bullet), style="List Bullet")
        return

    if kind == "projects":
        line = document.add_paragraph()
        run = line.add_run(str(item.get("name") or "Project"))
        run.bold = True
        if item.get("role"):
            line.add_run(f" — {item['role']}")
        if item.get("description"):
            document.add_paragraph(str(item["description"]))
        techs = ", ".join(item.get("technologies") or [])
        if techs:
            document.add_paragraph(f"Technologies: {techs}")
        for bullet in item.get("bullets") or []:
            document.add_paragraph(str(bullet), style="List Bullet")
        return

    if kind == "certifications":
        bits = [item.get("name"), item.get("issuer"), item.get("date")]
        document.add_paragraph(" — ".join(part for part in bits if part))
        return

    title = item.get("title") or item.get("name") or ""
    if title:
        p = document.add_paragraph()
        p.add_run(str(title)).bold = True
        if item.get("subtitle"):
            p.add_run(f" — {item['subtitle']}")
    if item.get("description"):
        document.add_paragraph(str(item["description"]))
    for bullet in item.get("bullets") or []:
        document.add_paragraph(str(bullet), style="List Bullet")


def _section_has_content(section: dict[str, Any]) -> bool:
    kind = section.get("kind")
    if kind == "text":
        return bool(str(section.get("content") or "").strip())
    if kind == "skills":
        return any(g.get("skills") for g in section.get("groups") or [])
    if kind == "tags":
        return bool(section.get("tags"))
    return bool(section.get("items"))
