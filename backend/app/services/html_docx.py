"""Turn captured resume preview HTML into a Word document.

The editor sends the same markup used on screen, so Academic Simple (and every
other template) keeps its header, headings, columns and photo in the DOCX.
"""

from __future__ import annotations

import base64
import io
import re
from html.parser import HTMLParser
from typing import Any

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Mm, Pt, RGBColor

VOID = {
    "area",
    "base",
    "br",
    "col",
    "embed",
    "hr",
    "img",
    "input",
    "link",
    "meta",
    "param",
    "source",
    "track",
    "wbr",
}

SKIP_TAGS = {"script", "style"}

DOCX_FONTS = {
    "Inter": "Calibri",
    "Georgia": "Georgia",
    "Garamond": "Garamond",
    "Calibri": "Calibri",
    "Arial": "Arial",
    "Helvetica": "Arial",
    "Times New Roman": "Times New Roman",
    "Cambria": "Cambria",
    "Verdana": "Verdana",
    "Lato": "Calibri",
    "Source Sans 3": "Calibri",
    "IBM Plex Sans": "Calibri",
}

Node = dict[str, Any]


class _TreeParser(HTMLParser):
    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.root: Node = {"tag": "document", "attrs": {}, "children": []}
        self.stack = [self.root]
        self._skip = 0

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        if self._skip:
            if tag not in VOID:
                self._skip += 1
            return
        if tag in SKIP_TAGS:
            self._skip = 1
            return
        node: Node = {"tag": tag, "attrs": {key: value or "" for key, value in attrs}, "children": []}
        self.stack[-1]["children"].append(node)
        if tag not in VOID:
            self.stack.append(node)

    def handle_endtag(self, tag: str) -> None:
        if self._skip:
            self._skip = max(0, self._skip - 1)
            return
        if tag in VOID:
            return
        for index in range(len(self.stack) - 1, 0, -1):
            if self.stack[index]["tag"] == tag:
                del self.stack[index:]
                break

    def handle_data(self, data: str) -> None:
        if self._skip or not data:
            return
        if not data.strip():
            return
        self.stack[-1]["children"].append({"tag": "#text", "attrs": {}, "children": [], "text": data})


def parse_html_tree(html_document: str) -> Node:
    parser = _TreeParser()
    parser.feed(html_document)
    parser.close()
    return parser.root


def classes_of(node: Node) -> list[str]:
    return [part for part in str(node.get("attrs", {}).get("class") or "").split() if part]


def has_class(node: Node, name: str) -> bool:
    return name in classes_of(node)


def has_class_prefix(node: Node, prefix: str) -> bool:
    return any(item.startswith(prefix) for item in classes_of(node))


def text_of(node: Node) -> str:
    if node.get("tag") == "#text":
        return str(node.get("text") or "")
    parts: list[str] = []
    for child in node.get("children") or []:
        parts.append(text_of(child))
    return re.sub(r"\s+", " ", "".join(parts)).strip()


def find(node: Node, predicate) -> Node | None:
    if predicate(node):
        return node
    for child in node.get("children") or []:
        match = find(child, predicate)
        if match is not None:
            return match
    return None


def children_elements(node: Node) -> list[Node]:
    return [child for child in node.get("children") or [] if child.get("tag") != "#text"]


def style_map(node: Node) -> dict[str, str]:
    raw = str(node.get("attrs", {}).get("style") or "")
    result: dict[str, str] = {}
    for part in raw.split(";"):
        if ":" not in part:
            continue
        key, value = part.split(":", 1)
        result[key.strip()] = value.strip()
    return result


def css_vars(node: Node) -> dict[str, str]:
    return {key: value for key, value in style_map(node).items() if key.startswith("--")}


def parse_color(value: str | None, fallback: tuple[int, int, int] = (17, 24, 39)) -> RGBColor:
    if not value:
        return RGBColor(*fallback)
    text = value.strip().lower()
    match = re.search(r"#([0-9a-f]{3,8})", text)
    if match:
        hex_value = match.group(1)
        if len(hex_value) == 3:
            hex_value = "".join(ch * 2 for ch in hex_value)
        hex_value = hex_value[:6]
        return RGBColor(int(hex_value[0:2], 16), int(hex_value[2:4], 16), int(hex_value[4:6], 16))
    rgb = re.search(r"rgba?\(\s*(\d+)\s*,\s*(\d+)\s*,\s*(\d+)", text)
    if rgb:
        return RGBColor(int(rgb.group(1)), int(rgb.group(2)), int(rgb.group(3)))
    return RGBColor(*fallback)


def first_font(family: str | None) -> str:
    if not family:
        return "Calibri"
    token = family.split(",")[0].strip().strip("\"'")
    return DOCX_FONTS.get(token, token or "Calibri")


def parse_pt(value: str | None, default: float) -> float:
    if not value:
        return default
    match = re.search(r"([0-9.]+)", value)
    return float(match.group(1)) if match else default


class _DocxTheme:
    def __init__(self, settings: dict[str, Any], vars_: dict[str, str]) -> None:
        family = vars_.get("--resume-font-family") or settings.get("fontFamily") or "Calibri"
        self.font = first_font(str(family))
        self.size = parse_pt(
            vars_.get("--resume-font-size"),
            float(settings.get("fontSize") or 10.5),
        )
        self.accent = parse_color(vars_.get("--resume-accent") or settings.get("accentColor"), (17, 24, 39))
        self.ink = parse_color(vars_.get("--resume-ink"), (17, 17, 17))
        self.muted = parse_color(vars_.get("--resume-muted"), (75, 85, 99))
        self.line_height = float(vars_.get("--resume-line-height") or settings.get("lineHeight") or 1.4)
        self.photo_pt = parse_pt(vars_.get("--resume-photo-size"), 72)


def _set_run(
    run,
    theme: _DocxTheme,
    *,
    size: float | None = None,
    bold: bool = False,
    color: RGBColor | None = None,
    italic: bool = False,
) -> None:
    run.font.name = theme.font
    run.font.size = Pt(size or theme.size)
    run.bold = bold
    run.italic = italic
    run.font.color.rgb = color or theme.ink
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.append(rFonts)
    rFonts.set(qn("w:ascii"), theme.font)
    rFonts.set(qn("w:hAnsi"), theme.font)


def _spacing(paragraph, before: float = 0, after: float = 4, line: float = 1.4) -> None:
    fmt = paragraph.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line
    fmt.line_spacing_rule = WD_LINE_SPACING.MULTIPLE


def _border(paragraph, color: RGBColor, size: str = "6", space: str = "1") -> None:
    pPr = paragraph._p.get_or_add_pPr()
    pBdr = pPr.find(qn("w:pBdr"))
    if pBdr is None:
        pBdr = OxmlElement("w:pBdr")
        pPr.append(pBdr)
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), size)
    bottom.set(qn("w:space"), space)
    bottom.set(qn("w:color"), f"{color[0]:02X}{color[1]:02X}{color[2]:02X}")
    existing = pBdr.find(qn("w:bottom"))
    if existing is not None:
        pBdr.remove(existing)
    pBdr.append(bottom)


def _shade(paragraph, fill: RGBColor) -> None:
    pPr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), f"{fill[0]:02X}{fill[1]:02X}{fill[2]:02X}")
    pPr.append(shd)


def _cell_shade(cell, fill: RGBColor) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), f"{fill[0]:02X}{fill[1]:02X}{fill[2]:02X}")
    tcPr.append(shd)


def _set_table_borders_none(table) -> None:
    tblPr = table._tbl.tblPr
    if tblPr is None:
        tblPr = OxmlElement("w:tblPr")
        table._tbl.insert(0, tblPr)
    borders = tblPr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tblPr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        element = borders.find(qn(f"w:{edge}"))
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "nil")
        element.set(qn("w:sz"), "0")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), "auto")


def _add_picture(paragraph, src: str, size_pt: float) -> bool:
    if not src.startswith("data:image/") or "svg" in src[:40] or "," not in src:
        return False
    try:
        payload = base64.b64decode(src.split(",", 1)[1])
    except Exception:  # noqa: BLE001
        return False
    if not payload:
        return False
    run = paragraph.add_run()
    try:
        run.add_picture(io.BytesIO(payload), width=Pt(size_pt), height=Pt(size_pt))
    except Exception:  # noqa: BLE001
        return False
    return True


def apply_page_setup(document: Document, settings: dict[str, Any]) -> None:
    section = document.sections[0]
    page_size = str(settings.get("pageSize") or settings.get("page_size") or "a4").lower()
    if page_size == "letter":
        section.page_width = Inches(8.5)
        section.page_height = Inches(11)
    else:
        section.page_width = Mm(210)
        section.page_height = Mm(297)
    margin = {"narrow": 0.43, "normal": 0.63, "wide": 0.87}.get(str(settings.get("margin") or "normal"), 0.63)
    for edge in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
        setattr(section, edge, Inches(margin))


def docx_from_resume_html(html_document: str, settings: dict[str, Any] | None = None) -> Document | None:
    if not html_document or "resume-" not in html_document:
        return None
    tree = parse_html_tree(html_document)
    root = find(tree, lambda node: has_class(node, "resume-document") or has_class(node, "resume-page"))
    if root is None:
        return None

    settings = settings or {}
    vars_ = css_vars(find(tree, lambda node: has_class(node, "resume-document")) or root)
    theme = _DocxTheme(settings, vars_)
    document = Document()
    apply_page_setup(document, settings)
    ctx = {"theme": theme, "ink": theme.ink, "muted": theme.muted, "on_dark": False}
    _render_container(document, root, ctx)
    return document


def _render_container(document: Document, node: Node, ctx: dict[str, Any], *, cell=None) -> None:
    target = cell if cell is not None else document
    if has_class(node, "resume-page__boundary") or has_class(node, "resume-page__number"):
        return
    if node.get("tag") == "img":
        src = str(node.get("attrs", {}).get("src") or "")
        paragraph = target.add_paragraph()
        _spacing(paragraph, after=6, line=ctx["theme"].line_height)
        _add_picture(paragraph, src, ctx["theme"].photo_pt)
        return
    if has_class(node, "resume-columns"):
        _render_columns(document, node, ctx, cell=cell)
        return
    if node.get("tag") == "header" or has_class_prefix(node, "resume-header"):
        _render_header(target, node, ctx)
        return
    if has_class(node, "resume-section"):
        _render_section(target, node, ctx)
        return
    if has_class(node, "resume-entry"):
        _render_entry(target, node, ctx)
        return
    if node.get("tag") in {"ul", "ol"} or has_class(node, "resume-bullets"):
        _render_list(target, node, ctx)
        return
    if has_class(node, "resume-skills__row") or has_class(node, "resume-inline-list") or has_class(
        node, "resume-tag-list"
    ) or has_class(node, "resume-contact-grid"):
        text = text_of(node)
        if text:
            paragraph = target.add_paragraph()
            _spacing(paragraph, after=2, line=ctx["theme"].line_height)
            _set_run(paragraph.add_run(text), ctx["theme"], color=ctx["ink"])
        return
    if node.get("tag") == "br":
        return
    if node.get("tag") == "p":
        text = text_of(node)
        if text:
            paragraph = target.add_paragraph()
            _spacing(paragraph, after=4, line=ctx["theme"].line_height)
            _set_run(paragraph.add_run(text), ctx["theme"], color=ctx["ink"])
        return
    elements = children_elements(node)
    if not elements:
        text = text_of(node)
        if text:
            paragraph = target.add_paragraph()
            _spacing(paragraph, after=2, line=ctx["theme"].line_height)
            _set_run(paragraph.add_run(text), ctx["theme"], color=ctx["ink"])
        return
    for child in elements:
        _render_container(document, child, ctx, cell=cell)


def _render_header(target, node: Node, ctx: dict[str, Any]) -> None:
    theme: _DocxTheme = ctx["theme"]
    empty = {"tag": "#text", "text": ""}
    name = text_of(find(node, lambda item: has_class(item, "resume-name")) or empty)
    role = text_of(find(node, lambda item: has_class(item, "resume-role")) or empty)
    contact = text_of(find(node, lambda item: has_class(item, "resume-contact")) or empty)
    photo = find(node, lambda item: item.get("tag") == "img")
    centered = has_class(node, "text-center") or has_class(node, "resume-header--photo-above")
    align = WD_ALIGN_PARAGRAPH.CENTER if centered else WD_ALIGN_PARAGRAPH.LEFT
    styles = style_map(node)
    name_color = RGBColor(255, 255, 255) if ctx["on_dark"] else theme.accent

    if photo is not None:
        src = str(photo.get("attrs", {}).get("src") or "")
        picture = target.add_paragraph()
        picture.alignment = align
        _spacing(picture, after=4, line=theme.line_height)
        _add_picture(picture, src, theme.photo_pt)

    if name:
        heading = target.add_paragraph()
        heading.alignment = align
        _spacing(heading, after=2, line=theme.line_height)
        _set_run(heading.add_run(name), theme, size=theme.size * 2, bold=True, color=name_color)
        if "border-bottom" in styles:
            _border(heading, theme.accent, size="8")
    if role:
        line = target.add_paragraph()
        line.alignment = align
        _spacing(line, after=2, line=theme.line_height)
        _set_run(line.add_run(role), theme, size=theme.size * 1.1, italic=True, color=ctx["muted"])
    if contact:
        line = target.add_paragraph()
        line.alignment = align
        _spacing(line, after=8, line=theme.line_height)
        _set_run(line.add_run(contact.replace(" | ", "  ·  ")), theme, size=theme.size * 0.92, color=ctx["muted"])
        if "border-bottom" in styles and not name:
            _border(line, theme.accent, size="8")


def _render_section(target, node: Node, ctx: dict[str, Any]) -> None:
    theme: _DocxTheme = ctx["theme"]
    title_node = find(node, lambda item: has_class(item, "resume-section__title"))
    title = text_of(title_node) if title_node else ""
    if title:
        heading = target.add_paragraph()
        _spacing(heading, before=10, after=3, line=theme.line_height)
        color = RGBColor(255, 255, 255) if ctx["on_dark"] else theme.accent
        class_names = classes_of(title_node or {})
        pill = "resume-section__title--pill" in class_names
        boxed = "background" in style_map(title_node or {})
        if pill or boxed:
            _shade(heading, theme.accent)
            color = RGBColor(255, 255, 255)
        _set_run(heading.add_run(title), theme, size=theme.size * 1.05, bold=True, color=color)
        if "resume-section__title--underline-accent" in class_names:
            _border(heading, theme.accent, size="12", space="2")
        elif "resume-section__title--inline-rule" in class_names or find(
            node, lambda item: has_class(item, "resume-section__rule")
        ):
            rule_color = RGBColor(255, 255, 255) if ctx["on_dark"] else theme.accent
            _border(heading, rule_color, size="6")
        elif "resume-section__title--gradient-bar" in class_names:
            pPr = heading._p.get_or_add_pPr()
            pBdr = OxmlElement("w:pBdr")
            left = OxmlElement("w:left")
            left.set(qn("w:val"), "single")
            left.set(qn("w:sz"), "18")
            left.set(qn("w:space"), "4")
            left.set(qn("w:color"), f"{theme.accent[0]:02X}{theme.accent[1]:02X}{theme.accent[2]:02X}")
            pBdr.append(left)
            pPr.append(pBdr)

    for child in children_elements(node):
        if (
            has_class(child, "resume-section__title")
            or child.get("tag") in {"h1", "h2", "h3"}
            or has_class(child, "resume-section__rule")
        ):
            continue
        nested_cell = target if hasattr(target, "_tc") else None
        _render_container(target, child, ctx, cell=nested_cell)


def _render_entry(target, node: Node, ctx: dict[str, Any]) -> None:
    theme: _DocxTheme = ctx["theme"]
    empty = {"tag": "#text", "text": ""}
    title = text_of(find(node, lambda item: has_class(item, "resume-entry__title")) or empty)
    meta = text_of(find(node, lambda item: has_class(item, "resume-entry__meta")) or empty)
    subtitle = text_of(find(node, lambda item: has_class(item, "resume-entry__subtitle")) or empty)
    if title or meta:
        line = target.add_paragraph()
        _spacing(line, before=4, after=0, line=theme.line_height)
        if title:
            _set_run(line.add_run(title), theme, bold=True, color=ctx["ink"])
        if meta:
            _set_run(line.add_run("    "), theme, color=ctx["muted"])
            _set_run(line.add_run(meta), theme, size=theme.size * 0.92, color=ctx["muted"])
    if subtitle:
        line = target.add_paragraph()
        _spacing(line, after=1, line=theme.line_height)
        _set_run(line.add_run(subtitle), theme, color=ctx["muted"])

    for child in children_elements(node):
        if has_class(child, "resume-entry__head") or has_class(child, "resume-entry__title") or has_class(
            child, "resume-entry__meta"
        ) or has_class(child, "resume-entry__subtitle"):
            continue
        if child.get("tag") in {"ul", "ol"} or has_class(child, "resume-bullets"):
            _render_list(target, child, ctx)
            continue
        text = text_of(child)
        if text:
            paragraph = target.add_paragraph()
            _spacing(paragraph, after=1, line=theme.line_height)
            _set_run(paragraph.add_run(text), theme, color=ctx["ink"])


def _render_list(target, node: Node, ctx: dict[str, Any]) -> None:
    theme: _DocxTheme = ctx["theme"]
    items = [child for child in children_elements(node) if child.get("tag") == "li"]
    for item in items:
        text = text_of(item)
        if not text:
            continue
        paragraph = target.add_paragraph(style="List Bullet")
        _spacing(paragraph, after=1, line=theme.line_height)
        if paragraph.runs:
            paragraph.runs[0].text = text
            _set_run(paragraph.runs[0], theme, color=ctx["ink"])
        else:
            _set_run(paragraph.add_run(text), theme, color=ctx["ink"])


def _sidebar_fill(node: Node) -> RGBColor | None:
    classes = classes_of(node)
    if "resume-sidebar--dark" in classes:
        return RGBColor(31, 41, 55)
    if "resume-sidebar--dark-tinted" in classes:
        return RGBColor(55, 65, 81)
    if "resume-sidebar--tinted" in classes:
        return RGBColor(244, 245, 247)
    return None


def _render_columns(document: Document, node: Node, ctx: dict[str, Any], *, cell=None) -> None:
    columns = children_elements(node)
    if len(columns) < 2:
        for child in columns:
            _render_container(document, child, ctx, cell=cell)
        return
    host = cell if cell is not None else document
    table = host.add_table(rows=1, cols=2)
    _set_table_borders_none(table)
    table.autofit = True
    left_cell, right_cell = table.rows[0].cells
    _fill_column_cell(left_cell, columns[0], ctx)
    _fill_column_cell(right_cell, columns[1], ctx)


def _fill_column_cell(cell, node: Node, ctx: dict[str, Any]) -> None:
    fill = _sidebar_fill(node)
    on_dark = has_class(node, "resume-sidebar--dark") or has_class(node, "resume-sidebar--dark-tinted")
    child_ctx = dict(ctx)
    if on_dark:
        child_ctx["on_dark"] = True
        child_ctx["ink"] = RGBColor(255, 255, 255)
        child_ctx["muted"] = RGBColor(209, 213, 219)
    if fill is not None:
        _cell_shade(cell, fill)
    _render_container(cell, node, child_ctx, cell=cell)
